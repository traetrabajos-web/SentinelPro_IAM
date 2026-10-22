# -*- coding: utf-8 -*-
"""
Generador del Documento Word (.docx) Profesional:
Plan Estratégico de Evolución Tecnológica, Recomendaciones de Ingeniería y Roadmap
para SentinelPro Access & IAM Enterprise Suite.
"""
import os
import sys
import datetime
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
DOCS_DIR = os.path.join(BASE_DIR, 'docs')
OUT_DOCX_ROOT = os.path.join(BASE_DIR, 'Documentacion_Plan_Estrategico_Y_Recomendaciones.docx')
OUT_DOCX_DOCS = os.path.join(DOCS_DIR, 'Documentacion_Plan_Estrategico_Y_Recomendaciones.docx')

# Paleta Corporativa Profesional
AZUL_OSCURO = RGBColor(0x0F, 0x17, 0x2A)       # #0F172A
INDIGO_CORP = RGBColor(0x37, 0x30, 0xA3)       # #3730A3
AZUL_MEDIO = RGBColor(0x1E, 0x40, 0xAF)        # #1E40AF
FONDO_TABLA_HEADER = RGBColor(0x1E, 0x29, 0x3B) # #1E293B
FONDO_FILA_ALT = RGBColor(0xF8, 0xFA, 0xFC)    # #F8FAFC
FONDO_CALLOUT = RGBColor(0xF1, 0xF5, 0xF9)     # #F1F5F9
TEXTO_GRIS = RGBColor(0x33, 0x41, 0x55)        # #334155
TEXTO_MUTED = RGBColor(0x64, 0x74, 0x8B)       # #64748B
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
VERDE_EXITO = RGBColor(0x05, 0x96, 0x69)
AMBAR_ALERTA = RGBColor(0xD9, 0x77, 0x06)

def set_cell_bg(cell, rgb: RGBColor):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    hex_color = '{:02X}{:02X}{:02X}'.format(rgb[0], rgb[1], rgb[2])
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def set_cell_borders(cell, color='CBD5E1', sz=4):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        b = OxmlElement(f'w:{side}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), str(sz))
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), color)
        tcBorders.append(b)
    tcPr.append(tcBorders)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in (('top', top), ('bottom', bottom), ('left', left), ('right', right)):
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_field(p, field_text):
    run = p.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = field_text
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    r = run._r
    r.append(fldChar1)
    r.append(instrText)
    r.append(fldChar2)
    r.append(fldChar3)
    run.font.size = Pt(8.5)
    run.font.color.rgb = TEXTO_MUTED

def configure_header_footer(doc):
    for section in doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.3)
        section.right_margin = Cm(2.3)

        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("SentinelPro Access & IAM — Plan Estratégico y Recomendaciones de Ingeniería")
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = TEXTO_MUTED

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = fp.add_run("SentinelPro Enterprise  |  Plan de Evolución Tecnológica y Roadmap  |  Página ")
        r1.font.size = Pt(8.5)
        r1.font.color.rgb = TEXTO_MUTED
        add_field(fp, "PAGE")
        r2 = fp.add_run(" de ")
        r2.font.size = Pt(8.5)
        r2.font.color.rgb = TEXTO_MUTED
        add_field(fp, "NUMPAGES")

def add_hr(doc, color='3730A3', sz=12):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(sz))
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)

def add_h1(doc, text):
    p = doc.add_heading(level=1)
    run = p.add_run(text)
    run.font.name = 'Segoe UI'
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = AZUL_OSCURO
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(4)
    add_hr(doc, '3730A3', 14)

def add_h2(doc, text):
    p = doc.add_heading(level=2)
    run = p.add_run(text)
    run.font.name = 'Segoe UI'
    run.font.size = Pt(12.5)
    run.font.bold = True
    run.font.color.rgb = INDIGO_CORP
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)

def add_h3(doc, text):
    p = doc.add_heading(level=3)
    run = p.add_run(text)
    run.font.name = 'Segoe UI'
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = AZUL_MEDIO
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)

def add_p(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        rb = p.add_run(bold_prefix)
        rb.font.name = 'Segoe UI'
        rb.font.size = Pt(9.5)
        rb.font.bold = True
        rb.font.color.rgb = AZUL_OSCURO
    r = p.add_run(text)
    r.font.name = 'Segoe UI'
    r.font.size = Pt(9.5)
    r.font.color.rgb = TEXTO_GRIS
    return p

def add_bullet(doc, bold_lead, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    rb = p.add_run(bold_lead)
    rb.font.name = 'Segoe UI'
    rb.font.size = Pt(9)
    rb.font.bold = True
    rb.font.color.rgb = AZUL_OSCURO
    r = p.add_run(text)
    r.font.name = 'Segoe UI'
    r.font.size = Pt(9)
    r.font.color.rgb = TEXTO_GRIS

def add_callout(doc, title, text, border_color='3730A3', bg_color=FONDO_CALLOUT):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_bg(cell, bg_color)
    
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '24')
    left.set(qn('w:space'), '0')
    left.set(qn('w:color'), border_color)
    tcBorders.append(left)
    
    for side in ('top', 'bottom', 'right'):
        b = OxmlElement(f'w:{side}')
        b.set(qn('w:val'), 'none')
        tcBorders.append(b)
    tcPr.append(tcBorders)
    set_cell_margins(cell, 90, 90, 140, 140)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    rt = p.add_run(f"📌 {title}\n")
    rt.font.name = 'Segoe UI'
    rt.font.size = Pt(9.5)
    rt.font.bold = True
    rt.font.color.rgb = AZUL_OSCURO

    rc = p.add_run(text)
    rc.font.name = 'Segoe UI'
    rc.font.size = Pt(9)
    rc.font.color.rgb = TEXTO_GRIS

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_styled_table(doc, headers, rows):
    tbl = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr_cells = tbl.rows[0].cells
    for i, h_text in enumerate(headers):
        set_cell_bg(hdr_cells[i], FONDO_TABLA_HEADER)
        set_cell_borders(hdr_cells[i], '334155', 6)
        set_cell_margins(hdr_cells[i], 90, 90, 110, 110)
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(str(h_text))
        run.font.name = 'Segoe UI'
        run.font.size = Pt(9)
        run.font.bold = True
        run.font.color.rgb = BLANCO

    for r_idx, row_data in enumerate(rows):
        row_cells = tbl.rows[r_idx + 1].cells
        bg = FONDO_FILA_ALT if r_idx % 2 == 1 else BLANCO
        for c_idx, val in enumerate(row_data):
            set_cell_bg(row_cells[c_idx], bg)
            set_cell_borders(row_cells[c_idx], 'CBD5E1', 4)
            set_cell_margins(row_cells[c_idx], 80, 80, 100, 100)
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(val))
            run.font.name = 'Segoe UI'
            run.font.size = Pt(8.5)
            if c_idx == 0:
                run.font.bold = True
                run.font.color.rgb = AZUL_OSCURO
            else:
                run.font.color.rgb = TEXTO_GRIS

    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def build_recommendations_document():
    doc = Document()
    configure_header_footer(doc)

    # =========================================================================
    # PORTADA
    # =========================================================================
    p_top = doc.add_paragraph()
    p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_top.paragraph_format.space_before = Pt(36)
    r_badge = p_top.add_run("PLAN ESTRATÉGICO DE INGENIERÍA DE SOFTWARE & ROADMAP TECNOLÓGICO")
    r_badge.font.name = 'Segoe UI'
    r_badge.font.size = Pt(10)
    r_badge.font.bold = True
    r_badge.font.color.rgb = INDIGO_CORP

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(18)
    r_title = p_title.add_run("SENTINELPRO ACCESS & IAM\nPLAN MAESTRO DE EVOLUCIÓN TECNOLÓGICA Y RECOMENDACIONES DE ARQUITECTURA")
    r_title.font.name = 'Segoe UI'
    r_title.font.size = Pt(21)
    r_title.font.bold = True
    r_title.font.color.rgb = AZUL_OSCURO

    add_hr(doc, '3730A3', 20)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(6)
    p_sub.paragraph_format.space_after = Pt(20)
    r_sub = p_sub.add_run(
        "Diagnóstico de Arquitectura, Recomendaciones Técnicas Especializadas, Fortalecimiento de Seguridad Criptográfica, "
        "Migración a Base de Datos Relacional de Alto Desempeño, Integración de Hardware Físico en Portería (OSDP/RFID/LPR), "
        "Arquitectura Garita Offline-First, Microservicios/Eventos, DevOps, y Cronograma de Ejecución (30, 60, 90 y 180 Días)."
    )
    r_sub.font.name = 'Segoe UI'
    r_sub.font.size = Pt(10.5)
    r_sub.font.color.rgb = TEXTO_GRIS

    add_styled_table(doc, ["Parámetro", "Detalle del Documento"], [
        ["Sistema Auditado", "SentinelPro Access & IAM (Versión 2.0 / Node.js Backend)"],
        ["Propósito", "Guía ejecutiva y técnica de optimizaciones, mejores prácticas y nuevas funcionalidades"],
        ["Áreas Clave", "Persistencia RDBMS, Criptografía IAM, IoT/Hardware Garita, UI/UX, CI/CD, Alta Disponibilidad"],
        ["Alcance", "Transición de prototipo/piloto a infraestructura industrial de misión crítica (High Concurrency)"],
        ["Estado de QA Actual", "15 de 15 pruebas automatizadas aprobadas (100% éxito)"],
        ["Autoría", "Dirección de Arquitectura y Seguridad de Software — SentinelPro"],
        ["Fecha de Publicación", datetime.date.today().strftime("%d de %B de %Y")]
    ])

    doc.add_page_break()

    # =========================================================================
    # CAPÍTULO 1: DIAGNÓSTICO DEL ESTADO ACTUAL Y MATRIZ FODA
    # =========================================================================
    add_h1(doc, "1. Diagnóstico del Estado Actual y Evaluación de Madurez")
    
    add_p(doc, 
        "SentinelPro IAM ha alcanzado un hito clave de desarrollo: cuenta con un motor de autorización granular probado (RBAC + ALLOW/DENY), "
        "gestión de portería/garita con control de duplicidad de entradas, cálculo automático de jornadas y horas extras, control de aforo en tiempo real, "
        "gestión de visitantes con asignación de gafetes y trazabilidad completa de eventos en bitácoras de auditoría. Además, el backend opera de forma "
        "autocontenida utilizando librerías nativas de Node.js (node:http, node:crypto, node:fs), eliminando vulnerabilidades de dependencias infladas."
    )

    add_callout(doc, "Estado de Validación QA (Hito Alcanzado)", 
        "La suite de pruebas automatizadas (scripts/qa-test-suite.js) valida satisfactoriamente el 100% de los 15 casos de prueba críticos, "
        "incluyendo concurrencia de sesiones (HTTP 409 Conflict), tokens HMAC-SHA256, hash scrypt, bloqueo de credenciales suspendidas, "
        "reemplazo de sesión forzado e inmutabilidad de la auditoría.",
        border_color='059669'
    )

    add_h2(doc, "Matriz FODA Técnica del Proyecto")
    add_styled_table(doc, ["Dimensión", "Hallazgos y Factores Técnicos"], [
        ["Fortalezas (F)", "• Cero dependencias externas en tiempo de ejecución (bajo riesgo de Supply Chain Attacks).\n• Algoritmos criptográficos robustos (scrypt con sal de 16 bytes y clave derivada de 64 bytes, HMAC-SHA256 con timingSafeEqual).\n• Modelo de autorización híbrido probado con precedencia estricta en el Backend.\n• Frontend liviano y responsivo sin overhead de compiladores complejos."],
        ["Oportunidades (O)", "• Integración directa con torniquetes y talanqueras físicas mediante protocolos industriales (OSDP / Wiegand).\n• Comercialización como solución SaaS multitenant para parques industriales y edificios corporativos.\n• Conexión con sistemas ERP de nómina (SAP, Workday, Siigo) mediante API REST y Webhooks."],
        ["Debilidades (D)", "• Persistencia actual basada en archivo plano JSON (database/db.json) con operaciones síncronas de I/O que pueden bloquear el Event Loop bajo alta concurrencia.\n• Sesiones y tokens gestionados sin rotación de Refresh Tokens en cookies HttpOnly.\n• Ausencia de autenticación multifactor (MFA/TOTP) para cuentas administrativas.\n• Falta de capacidad offline en garita ante interrupciones de red local."],
        ["Amenazas (A)", "• Congestión y demoras en garita peatonal durante cambios de turno masivos (500+ trabajadores en 15 minutos).\n• Ataques de fuerza bruta o suplantación física si las credenciales QR no son dinámicas.\n• Pérdida o corrupción de datos si el servidor se apaga abruptamente durante la escritura en db.json."]
    ])

    # =========================================================================
    # CAPÍTULO 2: EJE 1 — BASE DE DATOS Y PERSISTENCIA DE ALTA CONCURRENCIA
    # =========================================================================
    add_h1(doc, "2. Eje 1 — Modernización de la Persistencia de Datos")

    add_p(doc, 
        "La persistencia actual en 'db.json' es idónea para entornos de demostración, pruebas unitarias y despliegues monousuario. "
        "Sin embargo, en una operación industrial con múltiples torniquetes registrando marcaciones simultáneamente en horas pico, "
        "el uso de JSON en disco presenta cuellos de botella por contención de I/O y riesgo de condiciones de carrera (Race Conditions). "
        "Se recomienda la transición hacia una base de datos relacional de grado empresarial."
    )

    add_h2(doc, "Recomendaciones Específicas de Persistencia")
    add_bullet(doc, "1. Migración a PostgreSQL 16 con Prisma ORM / Kysely: ", 
        "Aprovechar el esquema ya diseñado en 'database/schema.sql' y 'database/sentinelpro_iam.sql'. Implementar conexión mediante Connection Pooling (pg-pool con 20-50 conexiones concurrentes) y transacciones ACID reales con aislamiento READ COMMITTED.")
    
    add_bullet(doc, "2. Particionamiento Declarativo de Tablas de Gran Volumen: ", 
        "Las tablas 'registros_acceso' y 'audit_logs' acumulan millones de filas al año. Se recomienda particionarlas por rango mensual (RANGE PARTITION BY created_at), manteniendo los índices en memoria y garantizando consultas de reportes en menos de 20 ms.")
    
    add_bullet(doc, "3. Capa de Caché y Memoria con Redis 7: ", 
        "Almacenar en Redis la lista blanca de credenciales activas (DNI/RFID -> Estado: AUTORIZADO), permitiendo que la verificación en garita responda en <2 milisegundos sin tocar la base de datos primaria. Además, utilizar Redis para el almacenamiento de sesiones activas y lista negra de tokens revocados.")

    add_bullet(doc, "4. Estrategia de Respaldo Continuo (RPO < 5 min, RTO < 15 min): ", 
        "Configurar archivado continuo de WAL (Write-Ahead Logging) hacia almacenamiento S3/Cloud Storage, complementado con copias de seguridad lógicas diarias automatizadas.")

    add_styled_table(doc, ["Criterio de Evaluación", "Implementación Actual (db.json)", "Implementación Recomendada (PostgreSQL + Redis)"], [
        ["Capacidad de Concurrencia", "Baja (~10 marcaciones/seg)", "Muy Alta (5,000+ marcaciones/seg)"],
        ["Consistencia y Transacciones", "Emulada en memoria con writeFileSync", "Transacciones ACID nativas con Rollback automático"],
        ["Tiempo de Consulta de Historial", "Lineal O(N) filtrando arreglos JS", "O(log N) con índices B-Tree y BRIN sobre timestamps"],
        ["Riesgo de Corrupción", "Medio (apagón durante writeFileSync)", "Cero (protección WAL y replicación síncrona)"],
        ["Escalabilidad Horizontal", "No soportada (un solo nodo)", "Totalmente soportada (Read Replicas + Connection Pooler)"]
    ])

    # =========================================================================
    # CAPÍTULO 3: EJE 2 — FORTALECIMIENTO DE LA SEGURIDAD CRIPTOGRÁFICA E IAM
    # =========================================================================
    add_h1(doc, "3. Eje 2 — Fortalecimiento de la Seguridad Criptográfica e IAM")

    add_p(doc, 
        "La seguridad IAM es el núcleo de SentinelPro. Para alcanzar cumplimiento con estándares internacionales (ISO/IEC 27001, SOC 2 Tipo II, NIST SP 800-63B), "
        "se recomienda enriquecer el motor criptográfico con las siguientes capacidades:"
    )

    add_h2(doc, "Medidas de Seguridad Críticas a Implementar")
    add_bullet(doc, "1. Autenticación Multifactor (MFA / 2FA) basada en RFC 6238 (TOTP): ", 
        "Obligatoria para usuarios con roles 'SUPER_ADMIN', 'ADMIN_SEGURIDAD' y 'SUPERVISOR_RRHH'. Compatibilidad con Google Authenticator, Microsoft Authenticator y llaves físicas FIDO2/WebAuthn (YubiKey).")

    add_bullet(doc, "2. Arquitectura de Tokens OAuth2 / OIDC con Par de Tokens Rotativos: ", 
        "Reemplazar el token único actual por: a) Access Token JWT de corta vida (15 minutos), y b) Refresh Token criptográfico rotativo (vida de 7 días) almacenado exclusivamente en cookies HTTP-Only, Secure, SameSite=Strict, inmune a ataques XSS.")

    add_bullet(doc, "3. Integración con Directorios Corporativos (SSO / SAML 2.0 / OpenID Connect): ", 
        "Permitir que los colaboradores y administradores inicien sesión con las credenciales corporativas de Microsoft Entra ID (Azure AD), Google Workspace u Okta, facilitando el aprovisionamiento automático mediante SCIM 2.0.")

    add_bullet(doc, "4. Rate Limiting Adaptativo y Defensa contra Ataques de Fuerza Bruta: ", 
        "Implementar limitación de tasa basada en el algoritmo Token Bucket tanto por dirección IP como por identificador de cuenta. Bloquear temporalmente intentos de marcación reiterados con tarjetas inválidas o contraseñas erróneas (máximo 5 intentos en 1 minuto).")

    add_bullet(doc, "5. Hardening de Cabeceras HTTP y Protección CSRF: ", 
        "Integrar políticas estrictas de Content Security Policy (CSP), Strict-Transport-Security (HSTS), X-Content-Type-Options: nosniff, y tokens anti-CSRF para todas las mutaciones POST/PATCH/DELETE.")

    add_callout(doc, "Gestión de Secretos y Rotación Criptográfica", 
        "Actualmente AUTH_SECRET cuenta con un valor fallback en código. En producción es perentorio exigir que AUTH_SECRET "
        "sea inyectado exclusivamente vía variables de entorno o mediante un gestor de secretos (HashiCorp Vault, AWS Secrets Manager o GCP Secret Manager), "
        "con rotación programada cada 90 días sin desconectar sesiones en curso gracias a un esquema de doble clave (clave activa y clave anterior).",
        border_color='D97706'
    )

    # =========================================================================
    # CAPÍTULO 4: EJE 3 — INTEGRACIÓN DE HARDWARE FÍSICO Y OPERACIÓN DE GARITA
    # =========================================================================
    add_h1(doc, "4. Eje 3 — Integración de Hardware Físico en Portería (IoT & Edge)")

    add_p(doc, 
        "Para que SentinelPro IAM gobierne el mundo físico de forma automatizada, debe interconectarse con los dispositivos de campo instalados en la portería y garitas de acceso."
    )

    add_h2(doc, "Protocolos y Dispositivos de Campo")
    add_styled_table(doc, ["Tipo de Dispositivo", "Protocolo Recomendado", "Lógica de Integración y Beneficio Operativo"], [
        ["Torniquetes y Molinetes Peatonales", "OSDP v2 sobre RS-485 o Relés IP", "Envío de pulso de apertura de 500 ms únicamente tras validación positiva en Backend. Monitoreo del sensor de paso para detectar paso no completado o saltos."],
        ["Lectores RFID / NFC", "Mifare DESFire EV2/EV3 (Cifrado AES)", "Lectura segura del UID y sector cifrado de la tarjeta corporativa. Previene la clonación de tarjetas convencionales Mifare Classic de 13.56 MHz."],
        ["Terminales de Autogestión (Kioscos)", "WebSockets / WebUSB / HID", "Lector de código de barras / QR industrial láser conectado al terminal para lectura instantánea de cédulas y carnets digitales desde el móvil del trabajador."],
        ["Barreras Vehiculares de Garita", "Controladora IP Modbus / TCP", "Apertura automatizada de talanqueras para camiones y vehículos autorizados tras validar placa o tag RFID vehicular (UHF pasivo 900 MHz)."],
        ["Cámaras LPR / ANPR", "ONVIF Profile S/T + API REST", "Captura automática de la placa del vehículo en garita, cruzándola con el padrón de vehículos autorizados y vinculándola con el registro de acceso del chofer."]
    ])

    add_h2(doc, "Credenciales Móviles con Códigos QR Dinámicos (Antifraude)")
    add_p(doc, 
        "Una vulnerabilidad común en los códigos QR estáticos es que los trabajadores pueden compartir una captura de pantalla a un compañero para que fiche por ellos. "
        "Se recomienda implementar Códigos QR Dinámicos con TOTP: el código cambia cada 30 segundos en la aplicación móvil del trabajador, firmado criptográficamente con una clave simétrica única por empleado. "
        "El lector de garita valida la firma y el código en tiempo real, imposibilitando el uso de capturas de pantalla o fotografías."
    )

    # =========================================================================
    # CAPÍTULO 5: EJE 4 — ARQUITECTURA GARITA OFFLINE-FIRST (RESILIENCIA EN BORDE)
    # =========================================================================
    add_h1(doc, "5. Eje 4 — Arquitectura Garita Offline-First (Continuidad de Negocio)")

    add_p(doc, 
        "En un entorno industrial, la garita de portería no puede detenerse si ocurre una falla en el enlace de fibra óptica, corte del switch principal o caída del servidor central. "
        "El ingreso de cientos de trabajadores no debe verse interrumpido bajo ninguna circunstancia."
    )

    add_callout(doc, "Principio de Operación Offline-First", 
        "El terminal o tableta del vigilante debe estar dotado de un Agente Edge local con base de datos SQLite integrada. "
        "El agente almacena una réplica local compacta del padrón de trabajadores y zonas autorizadas. "
        "Si se pierde la conexión con el servidor central, el terminal conmuta automáticamente a MODO OFFLINE, permitiendo el registro ininterrumpido.",
        border_color='1E40AF'
    )

    add_h2(doc, "Mecanismo de Sincronización Bidireccional")
    add_bullet(doc, "1. Detección Automática de Conectividad (Heartbeat): ", 
        "Un ping cada 5 segundos hacia /api/health determina el estado Online/Offline del enlace de portería.")
    
    add_bullet(doc, "2. Cola de Eventos Locales Transaccionales: ", 
        "Cada marcación efectuada en modo offline se encola localmente en SQLite con UUID v4, timestamp de precisión microsegundo y firma local del vigilante.")
    
    add_bullet(doc, "3. Sincronización en Ráfaga con Resolución de Conflictos: ", 
        "Al reanudarse la red, el agente local envía los lotes en cola al endpoint /api/admin/acceso/sincronizar-offline. El servidor central procesa los eventos en orden cronológico real, actualiza el aforo y audita el origen con la etiqueta 'MODO_OFFLINE_RECUPERADO'.")

    # =========================================================================
    # CAPÍTULO 6: EJE 5 — ARQUITECTURA DEL BACKEND, CONCURRENCIA Y EVENTOS
    # =========================================================================
    add_h1(doc, "6. Eje 5 — Modernización del Backend y Arquitectura por Eventos")

    add_p(doc, 
        "El backend actual destaca por su limpieza y ausencia de dependencias. Para escalar a múltiples sedes geográficas y miles de usuarios concurrentes, "
        "se recomienda evolucionar la estructura hacia Clean Architecture y desacoplamiento por eventos:"
    )

    add_h2(doc, "Pilares de Evolución del Backend")
    add_bullet(doc, "1. Migración Gradual a TypeScript: ", 
        "Tipado estático exhaustivo de entidades (Trabajador, RegistroAcceso, PermisoEfectivo, Sesion, AuditoriaLog). Detección de errores en tiempo de compilación y mayor mantenibilidad del código base.")

    add_bullet(doc, "2. Event-Driven Architecture (EDA) con BullMQ / RabbitMQ: ", 
        "Desacoplar las tareas no críticas del hilo principal de la marcación en portería. Cuando un trabajador pasa por el torniquete, el backend responde en <50 ms confirmando el acceso; inmediatamente emite un evento 'AccesoRegistradoEvent' que procesa de forma asíncrona: a) Actualización de métricas en Redis, b) Verificación de alertas de seguridad, c) Envío de notificaciones a supervisores, d) Registro de auditoría.")

    add_bullet(doc, "3. Canales WebSockets / SSE en Vivo: ", 
        "Eliminar el sondeo periódico (polling). Utilizar Server-Sent Events (SSE) o WebSockets para emitir actualizaciones instantáneas hacia la consola del vigilante, el monitor de aforo y el panel del Director de Seguridad.")

    # =========================================================================
    # CAPÍTULO 7: EJE 6 — EXPERIENCIA DE USUARIO (UI/UX) Y MODO KIOSCO INDUSTRIAL
    # =========================================================================
    add_h1(doc, "7. Eje 6 — Experiencia de Usuario (UI/UX) y Modos de Operación")

    add_p(doc, 
        "La interfaz actual en 'public/app.js' y 'public/styles.css' es completa y funcional. Para optimizar el rendimiento en campo, se recomiendan los siguientes refinamientos ergonómicos y técnicos:"
    )

    add_h2(doc, "Mejoras de Interfaz de Usuario")
    add_styled_table(doc, ["Módulo / Vista", "Mejora Propuesta", "Impacto Operativo"], [
        ["Kiosco de Autoservicio", "Modo Fullscreen táctil con botones de 80x80px mínimo y paleta de alto contraste", "Permite marcaciones ágiles por parte de operarios con guantes o en condiciones de baja luminosidad."],
        ["Feedback Acústico y Háptico", "Emisión de tonos de audio diferenciados (Beep agudo: acceso concedido; Doble tono grave: acceso denegado)", "El vigilante escucha el resultado de la marcación a distancia sin necesidad de mirar permanentemente el monitor."],
        ["PWA para Dispositivos Móviles", "Instalación como Progressive Web App en tablets Android/iOS para rondines de seguridad", "Los guardias de seguridad pueden auditar personal en patios, almacenes y áreas remotas sin estar atados a la garita."],
        ["Diseño Accesible (WCAG 2.1 AA)", "Navegación completa por teclado (Tab/Enter), contraste mínimo de texto 4.5:1 y soporte para lectores de pantalla", "Cumplimiento normativo y facilidad de uso para operadores con diversas capacidades."],
        ["Visualización de Aforo en Planta", "Mapa interactivo SVG de la planta con indicadores de calor y capacidad por zona", "Detección inmediata de áreas sobrepobladas o cuellos de botella en emergencias y evacuaciones."]
    ])

    # =========================================================================
    # CAPÍTULO 8: EJE 7 — DEVOPS, CONTENEDORES, CI/CD Y OBSERVABILIDAD
    # =========================================================================
    add_h1(doc, "8. Eje 7 — DevOps, Infraestructura en la Nube y Observabilidad")

    add_p(doc, 
        "Un sistema de control de acceso físico requiere una disponibilidad del 99.99% (cuatro nueves). "
        "A continuación se detallan las prácticas DevOps e infraestructura recomendadas para garantizar alta disponibilidad y despliegues sin interrupción (Zero-Downtime):"
    )

    add_h2(doc, "Prácticas de Ingeniería DevOps")
    add_bullet(doc, "1. Contenedorización Multi-Stage con Docker: ", 
        "Crear un Dockerfile multi-etapa optimizado con imagen base Node.js Alpine o Distroless, reduciendo el tamaño del contenedor a menos de 100 MB y minimizando la superficie de vulnerabilidades de sistema operativo.")

    add_bullet(doc, "2. Orquestación Local y de Producción: ", 
        "Crear 'docker-compose.yml' que configure el ecosistema completo: servicio backend SentinelPro, base de datos PostgreSQL 16 con volumen persistente, instancia de Redis para caché y servidor Nginx como Reverse Proxy con terminación SSL/TLS.")

    add_bullet(doc, "3. Pipeline CI/CD Automatizado en GitHub Actions: ", 
        "Configurar un flujo de integración continua que ejecute en cada commit: a) Verificación de sintaxis y linter, b) Suite de pruebas automatizadas 'qa-test-suite.js', c) Análisis de seguridad estático SAST (Semgrep / Trivy), d) Construcción y publicación de imagen Docker.")

    add_bullet(doc, "4. Pila de Observabilidad y Monitoreo Proactivo: ", 
        "Reemplazar los console.log por un logger estructurado en formato JSON (Pino). Exponer endpoint /metrics para Prometheus y configurar dashboards en Grafana para vigilar: latencia de marcación en garita, tasa de errores 403/409, aforo global y consumo de memoria del servidor.")

    # =========================================================================
    # CAPÍTULO 9: CRONOGRAMA DE EJECUCIÓN (ROADMAP 30 - 60 - 90 - 180 DÍAS)
    # =========================================================================
    add_h1(doc, "9. Plan de Implementación Cronológico (Roadmap)")

    add_p(doc, 
        "Para llevar a cabo las mejoras de forma ordenada sin interrumpir las operaciones actuales, "
        "se establece un cronograma estructurado en 4 fases secuenciales:"
    )

    add_styled_table(doc, ["Fase y Plazo", "Objetivos Principales", "Entregables Clave", "Impacto"], [
        ["Fase 1: Estabilización y Persistencia\n(Días 1 a 30)", 
         "• Migración a PostgreSQL 16.\n• Conexión de Redis para sesiones y caché.\n• Corrección de suite QA al 100%.\n• Dockerfile y docker-compose.", 
         "• Esquema relacional activo.\n• 15/15 pruebas QA automatizadas en CI.\n• Contenedores listos para prueba.",
         "Alta Concurrencia y Resiliencia de Datos"],
        ["Fase 2: Seguridad Criptográfica e IAM\n(Días 31 a 60)", 
         "• Autenticación Multifactor (MFA/TOTP).\n• Tokens OAuth2 (Access + Refresh en HttpOnly).\n• Rate Limiting estricto.\n• Migración de módulos a TypeScript.", 
         "• 2FA activo para Administradores.\n• Mitigación total de ataques XSS/CSRF.\n• Código base tipado y robusto.",
         "Cumplimiento ISO 27001 y Cero Vulnerabilidades"],
        ["Fase 3: Hardware y Modo Offline\n(Días 61 a 90)", 
         "• Integración con hardware de torniquetes (OSDP/Wiegand).\n• Agente de Garita Offline-First con SQLite.\n• QR dinámicos anticlonación con TOTP.", 
         "• Kiosco operativo sin internet.\n• Apertura física de puertas validada.\n• Carnets móviles con QR dinámico.",
         "Continuidad Operativa al 100% en Garita"],
        ["Fase 4: Escala, SSO y Producción\n(Días 91 a 180)", 
         "• Integración SSO corporativo (Azure AD / Google).\n• LPR vehicular automatizado con cámaras.\n• Despliegue en clúster Kubernetes.\n• Monitoreo con Grafana/Prometheus.", 
         "• Acceso unificado empresarial.\n• Control vehicular manos libres.\n• Alta disponibilidad (99.99%).",
         "Solución Enterprise de Clase Mundial"]
    ])

    # =========================================================================
    # CAPÍTULO 10: MATRIZ DE RIESGOS Y PLAN DE MITIGACIÓN
    # =========================================================================
    add_h1(doc, "10. Matriz de Riesgos Técnicos y Contingencias")

    add_styled_table(doc, ["Riesgo Identificado", "Nivel", "Impacto Potencial", "Estrategia Preventiva y Plan de Contingencia"], [
        ["Caída del enlace de red en garita principal", "ALTO", "Imposibilidad de registrar entradas/salidas provocando aglomeraciones.", "Implementación de arquitectura Offline-First con SQLite local y sincronización asíncrona automática al restaurarse el enlace."],
        ["Intento de suplantación mediante foto de código QR", "MEDIO", "Ingreso no autorizado de terceros haciéndose pasar por empleados.", "Uso de Códigos QR Dinámicos (rotación cada 30 segundos mediante algoritmo TOTP) y verificación obligatoria de foto en monitor de garita."],
        ["Saturación por concurrencia en cambio de turno masivo", "ALTO", "Latencia superior a 3 segundos en el torniquete causando filas.", "Caché en memoria con Redis (lectura en <2 ms), Connection Pooling en PostgreSQL y separación de lectura/escritura."],
        ["Fuga o compromiso de credenciales de usuario administrador", "CRÍTICO", "Modificación no autorizada de privilegios y apertura de accesos restringidos.", "Activación obligatoria de MFA/2FA, bloqueo tras 5 intentos fallidos, cierre automático por inactividad y registro inmutable en bitácora."],
        ["Falla física en el electroimán o torniquete", "MEDIO", "Bloqueo mecánico de la vía de evacuación o acceso.", "Configuración 'Fail-Safe' (en ausencia de energía eléctrica el torniquete o electroimán libera el paso automáticamente) conforme a normativas de seguridad humana."]
    ])

    # =========================================================================
    # CAPÍTULO 11: CONCLUSIONES Y CHECKLIST DE ACCIONES INMEDIATAS
    # =========================================================================
    add_h1(doc, "11. Conclusiones y Checklist de Próximos Pasos Inmediatos")

    add_p(doc, 
        "SentinelPro Access & IAM se encuentra en una posición técnica inmejorable: su lógica de negocio está consolidada, su modelo de autorización "
        "es granular y flexible, y las pruebas de calidad demuestran un funcionamiento impecable bajo los escenarios de uso previstos. "
        "Siguiendo las recomendaciones detalladas en este plan, el proyecto evolucionará de una aplicación web robusta hacia una plataforma "
        "empresarial integral de control de acceso e identidad de nivel internacional."
    )

    add_h2(doc, "Checklist de Acciones Inmediatas (Sprints 1 y 2)")
    add_bullet(doc, "☑ Validación de Suite QA: ", "Completado — Suite ejecutada al 100% con 15/15 casos aprobados.")
    add_bullet(doc, "☐ Implementación de Dockerfile y docker-compose.yml: ", "Empaquetar el sistema para despliegues reproducibles con un solo comando.")
    add_bullet(doc, "☐ Creación de Pipeline GitHub Actions (.github/workflows/ci.yml): ", "Asegurar que cada pull request o commit valide automáticamente los 15 tests.")
    add_bullet(doc, "☐ Conexión de PostgreSQL con el script existente: ", "Migrar los datos iniciales de db.json hacia las tablas relacionales de database/schema.sql.")
    add_bullet(doc, "☐ Módulo de Códigos QR Dinámicos: ", "Integrar generación de QR temporales con refresco cada 30 segundos en el Kiosco y credenciales.")

    # Guardar en las dos rutas (raíz y docs)
    doc.save(OUT_DOCX_ROOT)
    doc.save(OUT_DOCX_DOCS)
    print(f"Documento Word generado exitosamente en:\n1. {OUT_DOCX_ROOT}\n2. {OUT_DOCX_DOCS}")

if __name__ == '__main__':
    build_recommendations_document()
